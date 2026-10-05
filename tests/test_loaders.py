import docx
import pytest

from notes_qa.chunking import Chunker
from notes_qa.loaders import UnsupportedFileError, iter_note_files, load_file


def test_docx_headings_become_sections(tmp_path):
    path = tmp_path / "dbms.docx"
    d = docx.Document()
    d.add_heading("Normalization", level=1)
    d.add_heading("Third Normal Form", level=2)
    d.add_paragraph("A relation is in 3NF if it is in 2NF and has no transitive dependency.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "1NF"
    table.rows[0].cells[1].text = "Atomic values"
    d.save(path)

    pages = load_file(path)
    assert pages[0].page is None
    chunks = Chunker(400, 0).chunk(pages)
    assert chunks[0].section == "Normalization > Third Normal Form"
    assert "transitive dependency" in chunks[0].text
    assert "1NF | Atomic values" in chunks[-1].text


def test_unsupported_and_directory_walk(tmp_path):
    (tmp_path / "a.md").write_text("# A\ntext")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "b.txt").write_text("more text")
    (tmp_path / "image.png").write_bytes(b"\x89PNG")
    assert [p.name for p in iter_note_files([tmp_path])] == ["a.md", "b.txt"]
    with pytest.raises(UnsupportedFileError):
        load_file(tmp_path / "image.png")
