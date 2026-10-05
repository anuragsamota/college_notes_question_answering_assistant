"""Load lecture notes from PDF, DOCX, Markdown and plain-text files.

Each loader yields ``Page`` objects so that page numbers survive all the way
through to the citations shown to the student.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}


@dataclass
class Page:
    source: str          # file name shown in citations
    text: str
    page: int | None     # 1-indexed page for PDFs, None otherwise


class UnsupportedFileError(ValueError):
    pass


def load_file(path: str | Path, display_name: str | None = None) -> list[Page]:
    path = Path(path)
    name = display_name or path.name
    ext = path.suffix.lower()
    if ext == ".pdf":
        return list(_load_pdf(path, name))
    if ext == ".docx":
        return [Page(name, _load_docx(path), None)]
    if ext in {".md", ".markdown", ".txt"}:
        return [Page(name, path.read_text(encoding="utf-8", errors="replace"), None)]
    raise UnsupportedFileError(f"Unsupported file type '{ext}' for {name}")


def iter_note_files(paths: Iterable[str | Path]) -> Iterator[Path]:
    """Expand directories recursively and keep only supported files."""
    for p in map(Path, paths):
        if p.is_dir():
            for child in sorted(p.rglob("*")):
                if child.is_file() and child.suffix.lower() in SUPPORTED_EXTENSIONS:
                    yield child
        elif p.is_file():
            yield p
        else:
            raise FileNotFoundError(p)


def _load_pdf(path: Path, name: str) -> Iterator[Page]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    for i, pdf_page in enumerate(reader.pages, start=1):
        text = _clean_pdf_text(pdf_page.extract_text() or "")
        if text.strip():
            yield Page(name, text, i)


def _clean_pdf_text(text: str) -> str:
    # Re-join words hyphenated across line breaks ("algo-\nrithm" -> "algorithm").
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    # Drop bare page-number lines that PDF extraction leaves behind.
    text = re.sub(r"(?m)^\s*(page\s*)?\d{1,4}\s*$", "", text, flags=re.IGNORECASE)
    return text


def _load_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        match = re.match(r"heading\s*(\d)", style)
        if match:
            # Turn Word headings into Markdown headings so the chunker sees sections.
            lines.append("#" * int(match.group(1)) + " " + text)
        elif style == "title":
            lines.append("# " + text)
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)
